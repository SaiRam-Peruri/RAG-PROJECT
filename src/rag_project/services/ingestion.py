"""
Document ingestion service — indexes Federal Contracting documents into ChromaDB.
Production-grade: proper logging, error handling, validation.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Dict, List, Tuple

from tqdm import tqdm

from ..config import settings
from ..core.chroma_client import get_chroma_manager
from ..core.utils import chunk_text
from ..core.validation import validate_file_path
from ..logging_config import get_logger

logger = get_logger("ingestion")

# ── Exclusion rules ────────────────────────────────────
DENY_PATH_RX = [
    re.compile(r".*[\\/]+02_Compliance_and_Security[\\/].*", re.IGNORECASE),
    re.compile(r".*[\\/]+Signed[\\/].*", re.IGNORECASE),
]
DENY_NAME_RX = [
    re.compile(r"^SIGN[-_]", re.IGNORECASE),
]
DENY_EXACT_FILENAMES = {"structure"}


def should_index(p: Path) -> bool:
    """Check if a file should be indexed."""
    if p.name in DENY_EXACT_FILENAMES:
        return False
    if p.suffix.lower() not in settings.allowed_extensions:
        return False
    s = str(p)
    if any(rx.match(s) for rx in DENY_PATH_RX):
        return False
    if any(rx.match(p.name) for rx in DENY_NAME_RX):
        return False
    return True


def extract_metadata(p: Path) -> Dict:
    """Extract structured metadata from file path."""
    parts = list(p.parts)
    fn_lower = p.name.lower()

    meta = {
        "source_path": str(p),
        "filename": p.name,
        "ext": p.suffix.lower(),
        "bucket": "other",
        "opportunity": "unknown",
        "authority": "unknown",
        "stage": "other",
        "doc_role": "general",
    }

    for anchor, bucket in [
        ("01_Active_Pursuits", "active"),
        ("02_Awarded_Contracts", "awarded"),
        ("03_Unsuccessful_Pursuits", "unsuccessful"),
        ("04_Archive", "archive"),
    ]:
        if anchor in parts:
            meta["bucket"] = bucket
            idx = parts.index(anchor)
            if idx + 1 < len(parts):
                meta["opportunity"] = parts[idx + 1]
            break

    if "01_Government_Issued" in parts:
        meta["authority"] = "government"
    elif any(d in parts for d in ("03_Proposal_History", "04_Proposal_Development", "02_Industry_Responses")):
        meta["authority"] = "vendor"

    if "Amendments_QA" in parts:
        meta["stage"] = "amendment_or_qa"
    elif "Final_Solicitations" in parts:
        meta["stage"] = "solicitation"
    elif "Draft_Solicitations" in parts:
        meta["stage"] = "context"
    elif "RFIs" in parts:
        meta["stage"] = "rfi"
    elif "Award_Documents" in parts:
        meta["stage"] = "award"
    elif any(d in parts for d in ("03_Proposal_History", "04_Proposal_Development")):
        meta["stage"] = "proposal"

    meta["doc_role"] = _detect_doc_role(fn_lower, parts, meta["authority"])
    return meta


def _detect_doc_role(filename: str, path_parts: List[str], authority: str) -> str:
    """Detect semantic role of document for section-specific retrieval."""
    if authority == "government":
        if any(kw in filename for kw in ("acceptability", )) or ("matrix" in filename and "technical" in filename):
            return "evaluation_criteria"
        if any(kw in filename for kw in ("section_m", "evaluation_factors", "rating_plan")):
            return "evaluation_criteria"
        if any(kw in filename for kw in ("section_l", "instructions", "proposal_prep")):
            return "instructions"
        if any(kw in filename for kw in ("sow", "pws", "statement_of_work", "performance_work")):
            return "technical_requirements"
        if any(kw in filename for kw in ("pricing", "clin", "price_schedule", "igce")):
            return "pricing"
        if any(kw in filename for kw in ("amendment", "qa", "question")):
            return "amendment_qa"
        if "rfp" in filename or "solicitation" in filename:
            return "general"

    elif authority == "vendor":
        if any(kw in filename for kw in ("past_performance", "pastperf", "pp_", "reference")):
            return "past_performance"
        if any(kw in filename for kw in ("technical_approach", "tech_approach", "solution")):
            return "technical"
        if any(kw in filename for kw in ("management", "mgmt", "org_chart", "staffing")):
            return "management"
        if any(kw in filename for kw in ("quality", "qa_plan", "qap")):
            return "quality"
        if any(kw in filename for kw in ("security", "ssp", "cybersecurity")):
            return "security"
        if any(kw in filename for kw in ("resume", "cv", "personnel", "bio")):
            return "personnel"

    return "general"


def choose_collection(meta: Dict) -> str:
    """Determine which collection a document belongs to."""
    if meta["authority"] == "government":
        if meta["stage"] == "award":
            return settings.coll_draft
        return settings.coll_auth
    return settings.coll_draft


# ── File loaders ───────────────────────────────────────

def _load_pdf(p: Path) -> List[Tuple[str, Dict]]:
    from pypdf import PdfReader
    out = []
    reader = PdfReader(str(p))
    for i, page in enumerate(reader.pages):
        txt = page.extract_text() or ""
        meta = extract_metadata(p)
        meta["page"] = i + 1
        out.append((txt, meta))
    return out


def _load_docx(p: Path) -> List[Tuple[str, Dict]]:
    from docx import Document
    doc = Document(str(p))
    txt = "\n".join(para.text for para in doc.paragraphs if para.text)
    meta = extract_metadata(p)
    meta["page"] = 1
    return [(txt, meta)]


def _load_xlsx(p: Path) -> List[Tuple[str, Dict]]:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True)
    out = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                rows.append(" | ".join(vals))
        txt = f"Sheet: {sheet.title}\n" + "\n".join(rows)
        meta = extract_metadata(p)
        meta.update({"sheet": sheet.title, "page": 1})
        out.append((txt, meta))
    return out


def load_file(p: Path) -> List[Tuple[str, Dict]]:
    """Load a single file, returning list of (text, metadata) pairs."""
    try:
        ext = p.suffix.lower()
        if ext == ".pdf":
            return _load_pdf(p)
        if ext == ".docx":
            return _load_docx(p)
        if ext == ".xlsx":
            return _load_xlsx(p)
        return []
    except Exception as e:
        logger.warning("Failed to load %s: %s: %s", p.name, type(e).__name__, e)
        return []


# ── Main ingestion ─────────────────────────────────────

def ingest(
    root_dir: Path | None = None,
    clean: bool = True,
    show_progress: bool = True,
) -> Dict[str, int]:
    """
    Ingest all documents from Federal_Contracting into ChromaDB.

    Args:
        root_dir: Override the default Federal_Contracting directory
        clean: Whether to delete existing collections first
        show_progress: Show tqdm progress bar

    Returns:
        Dict with stats: {"files": N, "chunks": N, "auth_chunks": N, "draft_chunks": N}
    """
    root = root_dir or settings.federal_contracting_dir
    if not root.exists():
        raise FileNotFoundError(f"Federal_Contracting directory not found: {root}")

    settings.require_api_key()
    manager = get_chroma_manager()

    if clean:
        manager.delete_collection(settings.coll_auth)
        manager.delete_collection(settings.coll_draft)
        logger.info("Cleared existing collections")

    coll_auth = manager.get_or_create_collection(settings.coll_auth)
    coll_draft = manager.get_or_create_collection(settings.coll_draft)

    files = [p for p in root.rglob("*") if p.is_file() and should_index(p)]
    logger.info("Found %d indexable files in %s", len(files), root)

    stats = {"files": len(files), "chunks": 0, "auth_chunks": 0, "draft_chunks": 0}
    iterator = tqdm(files, desc="Ingesting") if show_progress else files

    for p in iterator:
        loaded = load_file(p)
        for raw_text, meta in loaded:
            chunks = chunk_text(raw_text, settings.chunk_max_chars, settings.chunk_overlap)
            if not chunks:
                continue

            ids = [f"{meta['source_path']}::p{meta.get('page', 1)}::c{j}" for j, _ in enumerate(chunks)]
            target_name = choose_collection(meta)
            target = coll_auth if target_name == settings.coll_auth else coll_draft

            try:
                target.add(ids=ids, documents=chunks, metadatas=[meta] * len(chunks))
                stats["chunks"] += len(chunks)
                if target_name == settings.coll_auth:
                    stats["auth_chunks"] += len(chunks)
                else:
                    stats["draft_chunks"] += len(chunks)
            except Exception as e:
                logger.error("Failed to index %s: %s", p.name, e)

    logger.info(
        "Ingestion complete: %d files, %d chunks (auth=%d, draft=%d)",
        stats["files"], stats["chunks"], stats["auth_chunks"], stats["draft_chunks"],
    )
    return stats
