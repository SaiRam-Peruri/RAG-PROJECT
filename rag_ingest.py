import os
import re
from pathlib import Path
from typing import Dict, List, Tuple
from tqdm import tqdm

import chromadb
from chromadb.utils import embedding_functions

from pypdf import PdfReader
from docx import Document
import openpyxl


ROOT = Path(r"C:\Users\ACER\RAG-Project\Federal_Contracting")

# ---- Exclusion rules (hard boundaries) ----
DENY_PATH_RX = [
    re.compile(r".*[\\/]+02_Compliance_and_Security[\\/].*", re.IGNORECASE),
    re.compile(r".*[\\/]+Signed[\\/].*", re.IGNORECASE),
]
DENY_NAME_RX = [
    re.compile(r"^SIGN[-_]", re.IGNORECASE),
]
DENY_EXACT_FILENAMES = {"structure"}  # the helper file in repo root

ALLOWED_EXT = {".pdf", ".docx", ".xlsx"}

# ---- Two collections ----
COLL_AUTH = "authoritative"
COLL_DRAFT = "drafting"


def should_index(p: Path) -> bool:
    if p.name in DENY_EXACT_FILENAMES:
        return False
    if p.suffix.lower() not in ALLOWED_EXT:
        return False
    s = str(p)
    if any(rx.match(s) for rx in DENY_PATH_RX):
        return False
    if any(rx.match(p.name) for rx in DENY_NAME_RX):
        return False
    return True


def extract_metadata(p: Path) -> Dict:
    parts = list(p.parts)
    fn_lower = p.name.lower()
    
    meta = {
        "source_path": str(p),
        "filename": p.name,
        "ext": p.suffix.lower(),
        "bucket": "other",
        "opportunity": "unknown",
        "authority": "unknown",  # government | vendor
        "stage": "other",        # solicitation | amendment_or_qa | proposal | rfi | award | context
        "doc_role": "general",   # semantic role for section-specific retrieval
    }

    for anchor, bucket in [
        ("01_Active_Pursuits", "active"),
        ("02_Awarded_Contracts", "awarded"),
        ("03_Unsuccessful_Pursuits", "unsuccessful"),
        ("04_Archive", "archive"),
    ]:
        if anchor in parts:
            meta["bucket"] = bucket
            idx = parts.index(anchor)
            if idx + 1 < len(parts):
                meta["opportunity"] = parts[idx + 1]
            break

    if "01_Government_Issued" in parts:
        meta["authority"] = "government"
    elif "03_Proposal_History" in parts or "04_Proposal_Development" in parts or "02_Industry_Responses" in parts:
        meta["authority"] = "vendor"

    if "Amendments_QA" in parts:
        meta["stage"] = "amendment_or_qa"
    elif "Final_Solicitations" in parts:
        meta["stage"] = "solicitation"
    elif "Draft_Solicitations" in parts:
        meta["stage"] = "context"
    elif "RFIs" in parts:
        meta["stage"] = "rfi"
    elif "Award_Documents" in parts:
        meta["stage"] = "award"
    elif "03_Proposal_History" in parts or "04_Proposal_Development" in parts:
        meta["stage"] = "proposal"

    # Semantic role detection for intelligent retrieval
    meta["doc_role"] = detect_doc_role(fn_lower, parts, meta["authority"])

    return meta


def detect_doc_role(filename: str, path_parts: List[str], authority: str) -> str:
    """
    Detect semantic role of document for section-specific retrieval.
    Roles: evaluation_criteria, technical_requirements, instructions, 
           past_performance, management, pricing, quality, security
    """
    # Government docs
    if authority == "government":
        # RFP base files typically contain multiple sections - default to general
        # We'll tag amendments as amendment_qa but they'll still be searchable for evaluation
        is_rfp = "rfp" in filename or "solicitation" in filename
        is_amendment = "amendment" in filename
        
        # Technical Acceptability Matrix = evaluation criteria
        if "acceptability" in filename or ("matrix" in filename and "technical" in filename):
            return "evaluation_criteria"
        
        # Section M = Evaluation (explicit naming)
        if any(kw in filename for kw in ["section_m", "evaluation_factors", "rating_plan"]):
            return "evaluation_criteria"
        
        # Section L = Instructions (explicit naming)
        if any(kw in filename for kw in ["section_l", "instructions", "proposal_prep"]):
            return "instructions"
        
        # Technical requirements / SOW / PWS (explicit files)
        if any(kw in filename for kw in ["sow", "pws", "statement_of_work", "performance_work"]):
            return "technical_requirements"
        
        # Pricing / CLIN (explicit files)
        if any(kw in filename for kw in ["pricing", "clin", "price_schedule", "igce"]):
            return "pricing"
        
        # Amendments / Q&A - these contain mixed content
        # Tag as amendment_qa, but queries will include this when searching for evaluation
        if is_amendment or "qa" in filename or "question" in filename:
            return "amendment_qa"
        
        # Base RFP files (not specific sections) - these contain everything
        # Tag as general so they're available for broad searches
        if is_rfp:
            return "general"
    
    # Vendor/Internal docs
    elif authority == "vendor":
        # Past Performance
        if any(kw in filename for kw in ["past_performance", "pastperf", "pp_", "reference"]):
            return "past_performance"
        
        # Technical Approach
        if any(kw in filename for kw in ["technical_approach", "tech_approach", "solution"]):
            return "technical"
        
        # Management Plan
        if any(kw in filename for kw in ["management", "mgmt", "org_chart", "staffing"]):
            return "management"
        
        # Quality / QA
        if any(kw in filename for kw in ["quality", "qa_plan", "qap"]):
            return "quality"
        
        # Security
        if any(kw in filename for kw in ["security", "ssp", "cybersecurity"]):
            return "security"
        
        # Resumes / Key Personnel
        if any(kw in filename for kw in ["resume", "cv", "personnel", "bio"]):
            return "personnel"
    
    # Default
    return "general"


def chunk_text(text: str, max_chars: int = 3500, overlap: int = 300) -> List[str]:
    # Simple, robust chunker (char-based). Good enough to start.
    text = re.sub(r"\s+\n", "\n", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    n = len(text)
    while start < n:
        end = min(n, start + max_chars)
        chunk = text[start:end]
        chunks.append(chunk)
        if end == n:
            break
        start = max(0, end - overlap)
    return chunks


def load_pdf(p: Path) -> List[Tuple[str, Dict]]:
    out = []
    reader = PdfReader(str(p))
    for i, page in enumerate(reader.pages):
        txt = page.extract_text() or ""
        meta = extract_metadata(p)
        meta.update({"page": i + 1})
        out.append((txt, meta))
    return out


def load_docx(p: Path) -> List[Tuple[str, Dict]]:
    doc = Document(str(p))
    txt = "\n".join([para.text for para in doc.paragraphs if para.text])
    meta = extract_metadata(p)
    meta.update({"page": 1})  # docx doesn't have pages reliably
    return [(txt, meta)]


def load_xlsx(p: Path) -> List[Tuple[str, Dict]]:
    wb = openpyxl.load_workbook(str(p), data_only=True)
    out = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
            if vals:
                rows.append(" | ".join(vals))
        txt = f"Sheet: {sheet.title}\n" + "\n".join(rows)
        meta = extract_metadata(p)
        meta.update({"sheet": sheet.title, "page": 1})
        out.append((txt, meta))
    return out


def load_file(p: Path) -> List[Tuple[str, Dict]]:
    try:
        if p.suffix.lower() == ".pdf":
            return load_pdf(p)
        if p.suffix.lower() == ".docx":
            return load_docx(p)
        if p.suffix.lower() == ".xlsx":
            return load_xlsx(p)
        return []
    except Exception as e:
        print(f"\nWarning: Failed to load {p.name}: {type(e).__name__}: {e}")
        return []


def choose_collection(meta: Dict) -> str:
    # Authoritative index: only government-issued solicitation context (RFP + amendments/Q&A + attachments)
    # Drafting index: proposal history/development + optionally vendor industry responses
    # Award documents are excluded from authoritative to avoid confusion with pre-award solicitation requirements
    if meta["authority"] == "government":
        if meta["stage"] == "award":
            return COLL_DRAFT  # Award docs not authoritative for solicitation Q&A
        return COLL_AUTH
    return COLL_DRAFT


def main():
    if not os.getenv("OPENAI_API_KEY"):
        raise SystemExit("OPENAI_API_KEY not set. In PowerShell: $env:OPENAI_API_KEY='...'\n")

    client = chromadb.PersistentClient(path=str(Path("chroma_db").resolve()))

    embedder = embedding_functions.OpenAIEmbeddingFunction(
        api_key=os.environ["OPENAI_API_KEY"],
        model_name="text-embedding-3-large",
    )

    # Clear existing collections (idempotent re-run)
    # Comment these out if you want incremental indexing later.
    try:
        client.delete_collection(name=COLL_AUTH)
    except ValueError:
        pass  # Collection doesn't exist yet
    try:
        client.delete_collection(name=COLL_DRAFT)
    except ValueError:
        pass  # Collection doesn't exist yet

    coll_auth = client.get_or_create_collection(name=COLL_AUTH, embedding_function=embedder)
    coll_draft = client.get_or_create_collection(name=COLL_DRAFT, embedding_function=embedder)

    files = [p for p in ROOT.rglob("*") if p.is_file() and should_index(p)]
    print(f"Indexable files (after exclusions): {len(files)}")

    doc_count = 0

    for p in tqdm(files, desc="Ingesting"):
        loaded = load_file(p)
        for raw_text, meta in loaded:
            chunks = chunk_text(raw_text)
            if not chunks:
                continue

            ids = []
            docs = []
            metas = []

            for j, ch in enumerate(chunks):
                doc_id = f"{meta['source_path']}::p{meta.get('page',1)}::c{j}"
                ids.append(doc_id)
                docs.append(ch)
                metas.append(meta)

            target = choose_collection(meta)
            if target == COLL_AUTH:
                coll_auth.add(ids=ids, documents=docs, metadatas=metas)
            else:
                coll_draft.add(ids=ids, documents=docs, metadatas=metas)

            doc_count += len(chunks)

    print(f"Total chunks indexed: {doc_count}")
    print("Done. Collections:")
    print(f" - {COLL_AUTH} (government/authoritative)")
    print(f" - {COLL_DRAFT} (vendor/drafting)")


if __name__ == "__main__":
    main()
